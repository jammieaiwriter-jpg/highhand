# -*- coding: utf-8 -*-
"""h93_tools.py — 93H 進度追蹤：確定性工具箱（2026-09-11 第二批改善）

代理只做「判斷」；抓取、查找、寫回、記異動、推播、重產全部走這裡，
不要再在班次裡現場寫 openpyxl / urllib 腳本（每次重新發明＝每次可能寫錯）。
所有輸出 UTF-8 JSON 或純文字；行號一律＝Sheet 列號（第 1 列是標題）。

用法（在 進度追蹤/ 目錄下執行）：
  python h93_tools.py today [--out F]
        抓 ?today=1；tg 項目整理成 {date,time,sender,text}。印 JSON。
  python h93_tools.py snapshot [--out F]
        下載盤點 export 為基準 xlsx（預設 backup/雲端盤點_<時間戳>.xlsx），印路徑。
  python h93_tools.py dump <分頁> [--from F] [--tail N]
        列出分頁內容，r<列號>: 欄1 | 欄2 | …
  python h93_tools.py find <關鍵詞...> [--sheet S] [--from F]
        跨分頁搜尋含所有關鍵詞的列（查重用）。不給 --sheet 時掃 §8 的五個分頁。
  python h93_tools.py changelog [--date 9/11] [--from F]
        印異動紀錄（防重複入帳）。
  python h93_tools.py apply <變更單.json> [--base F] [--dry-run] [--run-id R] [--force-live]
        套用變更單 → 寫回雲端 → 重新下載驗證；自動補「異動紀錄」一列（來源尾加 run_id）。
        --base：代理先前讀的快照；若目標分頁在此期間被人改過，該分頁整頁放棄不寫。
        --dry-run：只算差異不寫。寫入模式檔 _state/write_mode.txt 為 dry 時，一律視為 --dry-run。
  python h93_tools.py pending list [--all]
  python h93_tools.py pending add --q "問題" [--ref "出工預核 r168"] [--run-id R]
  python h93_tools.py pending answer <編號> --text "回覆" [--source "TG 9/11 21:10"]
  python h93_tools.py pending close <編號> [--note "已寫入出工預核 r168"]
        「待確認」分頁＝人機回饋隊列（玲嬅也可直接在 Sheet 填回覆欄）。
  python h93_tools.py tg <訊息檔> [--dry-run]
        推 Telegram 給玲嬅（自動分段），印 message_id。
  python h93_tools.py rebuild [--no-push]
        重產儀表板並 push（與 自動產出上傳儀表板.bat 同邏輯），印 commit hash。
  python h93_tools.py mode
        印目前寫入模式（dry | live）。
  python h93_tools.py help
        印變更單格式。
"""
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.parse
import urllib.request
from datetime import date, datetime

BASE = os.path.dirname(os.path.abspath(__file__))
EXEC = ("https://script.google.com/macros/s/"
        "AKfycbxxBXks2VtVjBJZhkzkzTVexEpUmcyCKMKjE18-dpiEh2SnfHcwP2VG6SMPV1NfAblJjw/exec")
SSID = "1JBdrqfzQ_AF0fG8x_RT7wjZSTCjCfgdgnX9FRkQJSd0"
EXPORT = f"https://docs.google.com/spreadsheets/d/{SSID}/export?format=xlsx"
# 兩種執行位置（2026-09-21 起）：
#  - 桌機：本檔在 工地/進度追蹤/，工作檔放原地（backup/、_state/、_auto_tmp/）。
#  - 雲端 routine：本檔被 build_dashboard_entry.py 同步到 highhand repo 的 tools/；工作檔一律放系統暫存夾，
#    絕不落在 repo 裡（repo 是公開的 GitHub Pages，雲端班次最後會 git add -A）。
IN_REPO = os.path.basename(BASE) == "tools"
WORK = os.path.join(tempfile.gettempdir(), "h93") if IN_REPO else BASE
BACKUP_DIR = os.path.join(WORK, "backup")
REPO = os.path.dirname(BASE) if IN_REPO else os.path.join(BASE, "highhand_repo")
MODE_FILE = os.path.join(WORK, "_state", "write_mode.txt")   # 雲端沒有這個檔 → live
TG_CONFIG = os.path.join(BASE, "telegram", "config.json")    # 雲端沒有 → 改讀環境變數 H93_TG_TOKEN／H93_TG_CHAT

DEDUPE_SHEETS = ["出工預核", "行政時程", "廠商協調", "缺失改善", "發包與工作事項"]
CHANGELOG = "異動紀錄"
CHANGELOG_HDR = ["日期", "項目", "變更內容", "來源"]
PROTECTED = {"說明"}                      # 永不寫
PENDING = "待確認"
PENDING_HDR = ["提出日", "編號", "問題", "相關分頁/列", "提出班次", "狀態",
               "回覆", "回覆日", "回覆來源", "落實備註"]

HELP_EDITS = """變更單（JSON，UTF-8）格式：
{
  "run_id": "R20260911-2145-report",            # 可省略，用 --run-id
  "source": "TG轉送-Ling HuaWu(9/11日誌)",        # 異動紀錄「來源」欄；工具自動補「 [run_id]」
  "changelog": {"item": "9/11晚報入帳", "detail": "一句話摘要本批改了什麼"},
  "ops": [
    {"op": "set",    "sheet": "出工預核", "row": 162, "col": 5, "value": "有,2人",
     "expect": {"1": "9/11", "2": "毅文"}},                 # expect：該列指定欄必含子字串，否則此 op 失敗
    {"op": "note",   "sheet": "出工預核", "row": 134, "col": 6, "text": "9/11未出工",
     "expect": {"2": "冠維"}},                              # note：以「；」接在原值後
    {"op": "append", "sheet": "出工預核",
     "values": ["0115/9/11", "柯(油漆)", "戶內油漆(3人)", "9/11出工數#1", "有,3人", ""],
     "unless_exists": {"1": "0115/9/11", "2": "柯"}},       # 已有同日同廠商列 → 跳過不重複建
    {"op": "set_by_key",  "sheet": "門禁追蹤", "key_col": 2, "key_text": "D03", "col": 3, "value": "地腳鍊安裝中"},
    {"op": "note_by_key", "sheet": "石材追蹤", "key_col": 2, "key_text": "玄關門框", "col": 6, "text": "9/11 2F(1人)"}
  ]
}
規則：同一分頁只要有一個 op 失敗（expect 不符／key 找不到或多筆），該分頁整頁不寫，其他分頁照寫。
日期一律寫 "0115/9/11" 這種前置 0 的民國字串（避免 Sheets 轉成日期型別）。
「說明」分頁永不寫；MILESTONES 不在 Sheet 裡，別想從這裡改。"""


# ------------------------------------------------------------------ 基礎
def out(obj):
    print(json.dumps(obj, ensure_ascii=False, indent=1))


def roc_today(d=None):
    d = d or date.today()
    return f"0{d.year - 1911}/{d.month}/{d.day}"


def norm(v):
    if v is None:
        return ""
    if isinstance(v, float):
        return "%g" % v
    return str(v)


def http_get(url, timeout=120):
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    return urllib.request.urlopen(req, timeout=timeout).read()


def post_json(obj, timeout=180):
    """POST 到 exec。Apps Script 回 302 轉一次性網址；urllib 會自動跟 GET 拿到 body。
    拿不到 body 也不代表失敗（見 AGENTS §2），呼叫端要以重新下載驗證為準。"""
    body = json.dumps(obj, ensure_ascii=False).encode("utf-8")
    req = urllib.request.Request(EXEC, data=body, headers={"Content-Type": "application/json"})
    try:
        r = urllib.request.urlopen(req, timeout=timeout)
        return r.status, r.read().decode("utf-8", "replace")[:300]
    except urllib.error.HTTPError as e:
        return e.code, e.read().decode("utf-8", "replace")[:300]
    except Exception as e:  # noqa: BLE001
        return -1, repr(e)


def download(path=None):
    os.makedirs(BACKUP_DIR, exist_ok=True)
    path = path or os.path.join(BACKUP_DIR, f"雲端盤點_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    data = http_get(EXPORT)
    with open(path, "wb") as f:
        f.write(data)
    return path


def open_wb(path):
    from openpyxl import load_workbook
    return load_workbook(path, data_only=True)


def grid(wb, name):
    if name not in wb.sheetnames:
        return None
    ws = wb[name]
    return [[ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            for r in range(1, ws.max_row + 1)]


def ngrid(g):
    return [[norm(v) for v in row] for row in (g or [])]


def latest_snapshot():
    files = sorted(f for f in os.listdir(BACKUP_DIR) if f.startswith("雲端盤點_")) if os.path.isdir(BACKUP_DIR) else []
    return os.path.join(BACKUP_DIR, files[-1]) if files else None


def source_wb(args):
    """--from F 指定快照；否則現抓一份新的。"""
    p = argval(args, "--from")
    if not p:
        p = download()
        print(f"# 已下載最新快照 {os.path.relpath(p, BASE)}", file=sys.stderr)
    return open_wb(p), p


def argval(args, flag, default=None):
    if flag in args:
        i = args.index(flag)
        if i + 1 < len(args):
            return args[i + 1]
    return default


def write_mode():
    try:
        with open(MODE_FILE, encoding="utf-8-sig") as f:
            m = f.read().strip().lower()
            return m if m in ("dry", "live") else "dry"
    except FileNotFoundError:
        return "live"


_FEATURES = None


def features():
    """後端 ?ver=1 的 features；後端 v11 起含 vers/protected/shrink_guard。抓不到視為空。"""
    global _FEATURES
    if _FEATURES is None:
        try:
            _FEATURES = json.loads(http_get(EXEC + "?ver=1", timeout=30).decode("utf-8")).get("features") or []
        except Exception:  # noqa: BLE001
            _FEATURES = []
    return _FEATURES


def fetch_vers():
    """各分頁版本號 {分頁: int}；後端未支援回 None（工具退回純客端漂移比對）。"""
    if "vers" not in features():
        return None
    try:
        return json.loads(http_get(f"{EXEC}?vers=1&ssId={SSID}", timeout=30).decode("utf-8")).get("vers") or {}
    except Exception:  # noqa: BLE001
        return None


def import_sheet(name, values, expect_ver=None, create=False, allow_shrink=False):
    vals = [[("" if v is None else (v if isinstance(v, (int, float, str)) else str(v))) for v in row]
            for row in values]
    body = {"key": "93h", "type": "import_sheet", "ssId": SSID, "name": name, "clear": True, "values": vals}
    if expect_ver is not None:
        body["expectVer"] = expect_ver      # 後端 v11：版本不符回 {err:"stale"}，舊後端忽略此欄
    if create:
        body["create"] = True
    if allow_shrink:
        body["allowShrink"] = True
    return post_json(body)


def post_err(txt):
    """從 POST 回應抓 err 代碼（stale / shrink / protected / unknown sheet）；沒有回 None。"""
    try:
        j = json.loads(txt)
        return None if j.get("ok") else (j.get("err") or "unknown")
    except Exception:  # noqa: BLE001
        return None


def sheet_diff(want, got):
    """回傳 (差異列數, 前三筆差異)。兩邊先補齊欄寬。"""
    w = max([len(r) for r in want + got] or [0])
    a = [r + [""] * (w - len(r)) for r in want]
    b = [r + [""] * (w - len(r)) for r in got]
    diffs = []
    for i in range(max(len(a), len(b))):
        ra = a[i] if i < len(a) else None
        rb = b[i] if i < len(b) else None
        if ra != rb:
            diffs.append({"row": i + 1, "want": ra, "got": rb})
    return len(diffs), diffs[:3]


# ------------------------------------------------------------------ 指令：讀
def cmd_today(args):
    d = json.loads(http_get(EXEC + "?today=1", timeout=60).decode("utf-8"))

    def fix_tg(items):
        res = []
        for it in items or []:
            if isinstance(it, list) and len(it) >= 4:
                t = str(it[1])
                m = re.search(r"T(\d\d):(\d\d)", t)
                if m:  # Sheets 時間值以 1899-12-30 UTC 呈現，+8 為台北
                    hh = (int(m.group(1)) + 8) % 24
                    t = f"{hh:02d}:{m.group(2)}"
                res.append({"date": it[0], "time": t, "sender": it[2], "text": it[3]})
            else:
                res.append(it)
        return res

    d["tg"] = fix_tg(d.get("tg"))
    if isinstance(d.get("yday"), dict):
        d["yday"]["tg"] = fix_tg(d["yday"].get("tg"))
    d["_counts"] = {k: len(d.get(k) or []) for k in ("am", "pm", "sv", "raw", "tg")}
    d["_counts"]["yday"] = {k: len((d.get("yday") or {}).get(k) or []) for k in ("pm", "sv", "tg")}
    o = argval(args, "--out")
    if o:
        with open(o, "w", encoding="utf-8") as f:
            json.dump(d, f, ensure_ascii=False, indent=1)
        print(o)
    else:
        out(d)


def cmd_snapshot(args):
    # 一律印絕對路徑：雲端從 repo 根目錄執行、桌機從 進度追蹤/ 執行，相對路徑會因 cwd 不同而失效
    print(os.path.abspath(download(argval(args, "--out"))))


def cmd_dump(args):
    name = args[0]
    wb, _ = source_wb(args[1:])
    g = grid(wb, name)
    if g is None:
        print(f"沒有分頁「{name}」。有：{wb.sheetnames}")
        return
    tail = int(argval(args, "--tail", "0") or 0)
    start = max(1, len(g) - tail) if tail else 0
    if start > 0:
        print("r1: " + " | ".join(norm(v) for v in g[0]))
    for i in range(start, len(g)):
        print(f"r{i + 1}: " + " | ".join(norm(v) for v in g[i]))


def cmd_find(args):
    kws = [a for a in args if not a.startswith("--") and a != argval(args, "--from") and a != argval(args, "--sheet")]
    wb, _ = source_wb(args)
    sheets = [argval(args, "--sheet")] if argval(args, "--sheet") else DEDUPE_SHEETS
    hits = 0
    for s in sheets:
        g = grid(wb, s)
        if not g:
            continue
        for i, row in enumerate(g[1:], 2):
            txt = " ".join(norm(v) for v in row)
            if all(k in txt for k in kws):
                hits += 1
                print(f"[{s}] r{i}: " + " | ".join(norm(v) for v in row))
    print(f"# 共 {hits} 列命中 {kws} 於 {sheets}")


def cmd_changelog(args):
    wb, _ = source_wb(args)
    g = grid(wb, CHANGELOG) or []
    want = argval(args, "--date")
    for i, row in enumerate(g[1:], 2):
        d0 = norm(row[0]) if row else ""
        if want and not (d0 == want or d0.endswith("/" + want)):
            continue
        print(f"r{i}: " + " | ".join(norm(v) for v in row))


# ------------------------------------------------------------------ 指令：apply
def _find_rows(g, key_col, key_text):
    return [i for i, row in enumerate(g[1:], 2)
            if key_col - 1 < len(row) and key_text in norm(row[key_col - 1])]


def _check_expect(g, row, expect):
    if not expect:
        return None
    if row < 2 or row > len(g):
        return f"r{row} 不存在（分頁只有 {len(g)} 列）"
    for c, sub in expect.items():
        c = int(c)
        v = norm(g[row - 1][c - 1]) if c - 1 < len(g[row - 1]) else ""
        if sub not in v:
            return f"r{row} c{c} 期望含「{sub}」實為「{v}」"
    return None


def _ensure_width(g, col):
    for row in g:
        while len(row) < col:
            row.append(None)


def apply_ops(G, ops):
    """就地套用；回傳 {sheet: {"applied": n, "failed": [msg], "added": n}}"""
    rep = {}
    for k, op in enumerate(ops, 1):
        s = op.get("sheet")
        r = rep.setdefault(s, {"applied": 0, "failed": [], "added": 0, "skipped": []})
        g = G.get(s)
        if s in PROTECTED or g is None:
            r["failed"].append(f"op{k}: 分頁「{s}」不可寫或不存在")
            continue
        kind = op.get("op")
        try:
            if kind in ("set_by_key", "note_by_key"):
                rows = _find_rows(g, int(op["key_col"]), op["key_text"])
                if len(rows) != 1:
                    r["failed"].append(f"op{k}: key「{op['key_text']}」於 c{op['key_col']} 命中 {len(rows)} 列（需恰 1）")
                    continue
                op = dict(op, row=rows[0], op="set" if kind == "set_by_key" else "note")
                kind = op["op"]
            if kind in ("set", "note"):
                row, col = int(op["row"]), int(op["col"])
                err = _check_expect(g, row, op.get("expect"))
                if err:
                    r["failed"].append(f"op{k}: {err}")
                    continue
                if row < 2 or row > len(g):
                    r["failed"].append(f"op{k}: r{row} 不存在")
                    continue
                _ensure_width(g, col)
                if kind == "set":
                    g[row - 1][col - 1] = op["value"]
                else:
                    cur = norm(g[row - 1][col - 1])
                    g[row - 1][col - 1] = (cur + "；" + op["text"]) if cur else op["text"]
                r["applied"] += 1
            elif kind == "append":
                ue = op.get("unless_exists")
                if ue:
                    dup = [i for i, row in enumerate(g[1:], 2)
                           if all(sub in (norm(row[int(c) - 1]) if int(c) - 1 < len(row) else "")
                                  for c, sub in ue.items())]
                    if dup:
                        r["skipped"].append(f"op{k}: 已存在 r{dup[0]}，不重複新增")
                        continue
                g.append(list(op["values"]))
                r["applied"] += 1
                r["added"] += 1
            else:
                r["failed"].append(f"op{k}: 未知 op「{kind}」")
        except (KeyError, ValueError, TypeError) as e:
            r["failed"].append(f"op{k}: 欄位錯誤 {e!r}")
    return rep


def cmd_apply(args):
    path = args[0]
    with open(path, encoding="utf-8-sig") as f:
        edits = json.load(f)
    run_id = argval(args, "--run-id") or edits.get("run_id") or "manual"
    mode = write_mode()
    dry = "--dry-run" in args or (mode == "dry" and "--force-live" not in args)
    ops = edits.get("ops") or []
    touched = sorted({op.get("sheet") for op in ops} | {CHANGELOG})

    fresh = download()
    wb = open_wb(fresh)
    result = {"run_id": run_id, "write_mode": mode, "dry_run": dry, "fresh_snapshot": os.path.abspath(fresh),
              "sheets": {}, "warnings": []}

    # 並行編輯防護：代理讀的快照 vs 現在
    drift = set()
    basep = argval(args, "--base")
    if basep:
        bwb = open_wb(basep)
        for s in touched:
            if s == CHANGELOG:
                continue
            if ngrid(grid(bwb, s)) != ngrid(grid(wb, s)):
                drift.add(s)
                result["warnings"].append(f"「{s}」在你讀快照之後被改過，整頁放棄不寫")

    G = {s: grid(wb, s) for s in touched}
    if G.get(CHANGELOG) is None:
        G[CHANGELOG] = [CHANGELOG_HDR]
    before = {s: ngrid(G[s]) for s in touched}
    rep = apply_ops(G, [op for op in ops if op.get("sheet") not in drift])

    # 異動紀錄
    applied_any = any(v["applied"] for v in rep.values())
    if applied_any:
        cl = edits.get("changelog") or {}
        detail = cl.get("detail") or "；".join(f"{s}:{v['applied']}項" for s, v in rep.items() if v["applied"])
        src = f"{edits.get('source', '自動班次')} [{run_id}]"
        G[CHANGELOG].append([roc_today(), cl.get("item") or "自動入帳", detail, src])
        result["changelog_row"] = G[CHANGELOG][-1]

    # 決定要寫哪些分頁
    to_push = []
    for s in touched:
        r = rep.get(s, {"applied": 0, "failed": [], "added": 0, "skipped": []})
        entry = {"ops_applied": r["applied"], "ops_failed": r["failed"], "ops_skipped": r["skipped"],
                 "rows_added": r["added"]}
        after = ngrid(G[s])
        nd, sample = sheet_diff(after, before[s])
        entry["cells_rows_changed"] = nd
        entry["sample_changes"] = sample
        if s in drift:
            entry["status"] = "skipped_drift"
        elif r["failed"]:
            entry["status"] = "skipped_errors"
        elif nd == 0:
            entry["status"] = "unchanged"
        elif dry:
            entry["status"] = "dry"
        else:
            entry["status"] = "pending_push"
            to_push.append(s)
        result["sheets"][s] = entry
    if CHANGELOG in to_push and not applied_any:
        to_push.remove(CHANGELOG)

    if not dry:
        vers = fetch_vers()   # 後端 v11：拿 fresh 當下的版本號當 expectVer；舊後端為 None
        result["server_version_lock"] = vers is not None
        for s in to_push:
            st, txt = import_sheet(s, G[s], expect_ver=(vers or {}).get(s) if vers is not None else None)
            result["sheets"][s]["post"] = {"status": st, "body": txt}
            err = post_err(txt)
            if err:
                result["sheets"][s]["status"] = f"rejected_{err.replace(' ', '_')}"
                result["warnings"].append(f"「{s}」後端拒寫：{err}（{txt[:120]}）")
            time.sleep(1)
        to_push = [s for s in to_push if not result["sheets"][s]["status"].startswith("rejected_")]
        # 寫後驗證
        vpath = download(os.path.join(BACKUP_DIR, f"驗證_{run_id}.xlsx"))
        vwb = open_wb(vpath)
        total = 0
        for s in to_push:
            nd, sample = sheet_diff(ngrid(G[s]), ngrid(grid(vwb, s)))
            result["sheets"][s]["verify_diff"] = nd
            result["sheets"][s]["verify_sample"] = sample
            result["sheets"][s]["status"] = "pushed_ok" if nd == 0 else "pushed_MISMATCH"
            total += nd
        result["verify_total_diff"] = total
        try:
            os.remove(vpath)
        except OSError:
            pass

    result["summary"] = {s: v["status"] for s, v in result["sheets"].items()}
    rid_dir = os.path.join(WORK, "_auto_tmp", run_id)
    if os.path.isdir(rid_dir):
        with open(os.path.join(rid_dir, "apply_result.json"), "w", encoding="utf-8") as f:
            json.dump(result, f, ensure_ascii=False, indent=1)
    out(result)


# ------------------------------------------------------------------ 指令：pending
def _pending_grid(wb):
    g = grid(wb, PENDING)
    if not g:
        g = [PENDING_HDR]
    for row in g:
        while len(row) < len(PENDING_HDR):
            row.append(None)
    return g


def _pending_push(g, label, create=False):
    vers = fetch_vers()
    st, txt = import_sheet(PENDING, g, expect_ver=(vers or {}).get(PENDING) if vers is not None else None,
                           create=create)
    err = post_err(txt)
    if err:
        out({"action": label, "post": st, "rejected": err, "body": txt[:200]})
        sys.exit(1)
    vwb = open_wb(download(os.path.join(BACKUP_DIR, "驗證_pending.xlsx")))
    nd, _ = sheet_diff(ngrid(g), ngrid(grid(vwb, PENDING)))
    out({"action": label, "post": st, "verify_diff": nd, "rows": len(g) - 1})


def cmd_pending(args):
    sub = args[0] if args else "list"
    wb = open_wb(download())
    g = _pending_grid(wb)
    tab_missing = PENDING not in wb.sheetnames
    if sub == "list":
        show_all = "--all" in args
        for i, row in enumerate(g[1:], 2):
            if not show_all and norm(row[5]) not in ("待回覆", "已回覆"):
                continue
            print(f"r{i}: " + " | ".join(norm(v) for v in row))
        n_open = sum(1 for row in g[1:] if norm(row[5]) == "待回覆")
        n_ans = sum(1 for row in g[1:] if norm(row[5]) == "已回覆")
        print(f"# 待回覆 {n_open}、已回覆待落實 {n_ans}（共 {len(g) - 1} 筆）")
        return
    if sub == "add":
        q = argval(args, "--q")
        if not q:
            print("需要 --q 問題內容")
            sys.exit(2)
        nums = [int(norm(r[1])[1:]) for r in g[1:] if re.fullmatch(r"Q\d+", norm(r[1]))]
        qid = f"Q{(max(nums) + 1) if nums else 1:04d}"
        dup = [norm(r[1]) for r in g[1:] if norm(r[2]) == q and norm(r[5]) in ("待回覆", "已回覆")]
        if dup:
            out({"action": "add", "skipped": f"同一問題已在隊列 {dup[0]}"})
            return
        g.append([roc_today(), qid, q, argval(args, "--ref", ""), argval(args, "--run-id", ""),
                  "待回覆", "", "", "", ""])
        if write_mode() == "dry" and "--force-live" not in args:
            out({"action": "add", "dry_run": True, "would_add": g[-1]})
            return
        _pending_push(g, f"add {qid}", create=tab_missing)
        return
    if sub in ("answer", "close"):
        qid = args[1] if len(args) > 1 else ""
        rows = [i for i, r in enumerate(g[1:], 2) if norm(r[1]) == qid]
        if len(rows) != 1:
            print(f"編號「{qid}」命中 {len(rows)} 列")
            sys.exit(2)
        row = g[rows[0] - 1]
        if sub == "answer":
            row[6] = argval(args, "--text", "")
            row[7] = roc_today()
            row[8] = argval(args, "--source", "")
            row[5] = "已回覆"
        else:
            row[9] = argval(args, "--note", "")
            row[5] = "已落實" if "--cancel" not in args else "取消"
        if write_mode() == "dry" and "--force-live" not in args:
            out({"action": sub, "dry_run": True, "would_write": row})
            return
        _pending_push(g, f"{sub} {qid}")
        return
    print("pending 子指令：list|add|answer|close")
    sys.exit(2)


# ------------------------------------------------------------------ 指令：tg / rebuild / mode
def cmd_tg(args):
    with open(args[0], encoding="utf-8-sig") as f:
        msg = f.read().strip()
    chunks, cur = [], ""
    for line in msg.split("\n"):
        if len(cur) + len(line) + 1 > 3900:
            chunks.append(cur)
            cur = ""
        cur += line + "\n"
    if cur.strip():
        chunks.append(cur)
    if "--dry-run" in args:
        out({"dry_run": True, "chunks": len(chunks), "chars": len(msg), "head": msg[:200]})
        return
    if os.path.exists(TG_CONFIG):
        cfg = json.load(open(TG_CONFIG, encoding="utf-8"))
    else:   # 雲端 routine：金鑰不進 repo，由班次提示以環境變數帶入
        cfg = {"token": os.environ.get("H93_TG_TOKEN", ""), "chat_id": os.environ.get("H93_TG_CHAT", "")}
    if not cfg.get("token") or not cfg.get("chat_id"):
        out({"ok": False, "error": "沒有 Telegram 設定（桌機 telegram/config.json；雲端 H93_TG_TOKEN／H93_TG_CHAT）"})
        sys.exit(1)
    ids = []
    for c in chunks:
        data = urllib.parse.urlencode({"chat_id": cfg["chat_id"], "text": c}).encode("utf-8")
        r = urllib.request.urlopen(f"https://api.telegram.org/bot{cfg['token']}/sendMessage", data=data, timeout=60)
        d = json.loads(r.read().decode("utf-8"))
        ids.append(d.get("result", {}).get("message_id"))
        time.sleep(1)
    out({"ok": True, "message_ids": ids, "chunks": len(chunks)})


def cmd_rebuild(args):
    py = sys.executable
    if IN_REPO:   # 雲端：tools/cloud_build.py 產 93h/index.html；提交身分用 93H-cloud
        entry, cwd, msg = os.path.join(BASE, "cloud_build.py"), REPO, "cloud ledger update"
        ident = ["-c", "user.name=93H-cloud", "-c", "user.email=bot@93h"]
    else:
        entry, cwd, msg, ident = os.path.join(BASE, "build_dashboard_entry.py"), BASE, "auto update dashboard", []
    r = subprocess.run([py, entry], cwd=cwd,
                       capture_output=True, text=True, encoding="utf-8", errors="replace")
    res = {"build_rc": r.returncode, "build_tail": (r.stdout + r.stderr)[-600:]}
    if r.returncode != 0 or "--no-push" in args:
        res["pushed"] = False
        out(res)
        return

    def git(*a):
        return subprocess.run(["git", *a], cwd=REPO, capture_output=True, text=True,
                              encoding="utf-8", errors="replace")

    git("rebase", "--abort")
    git("add", "-A")
    if git("diff", "--cached", "--quiet").returncode == 0:
        res.update(pushed=False, note="NO CHANGE", head=git("rev-parse", "--short", "HEAD").stdout.strip())
        out(res)
        return
    git(*ident, "commit", "-m", msg, "--quiet")
    p = git("pull", "--rebase", "--quiet")
    if p.returncode != 0:
        git("checkout", "--theirs", "--", "93h/index.html", "93h/report/plan.json", "tools/builder.py")
        git("add", "-A")
        c = subprocess.run(["git", *ident, "rebase", "--continue"], cwd=REPO, capture_output=True, text=True,
                           env=dict(os.environ, GIT_EDITOR="true"))
        if c.returncode != 0:
            git("rebase", "--abort")
            res.update(pushed=False, error="REBASE FAILED")
            out(res)
            return
    ps = git("push", "--quiet")
    res.update(pushed=ps.returncode == 0, push_err=ps.stderr[-300:] if ps.returncode else "",
               head=git("rev-parse", "--short", "HEAD").stdout.strip())
    out(res)


def cmd_mode(args):
    print(write_mode())


# ------------------------------------------------------------------ main
CMDS = {"today": cmd_today, "snapshot": cmd_snapshot, "dump": cmd_dump, "find": cmd_find,
        "changelog": cmd_changelog, "apply": cmd_apply, "pending": cmd_pending, "tg": cmd_tg,
        "rebuild": cmd_rebuild, "mode": cmd_mode}


def main():
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:  # noqa: BLE001
        pass
    args = sys.argv[1:]
    if not args or args[0] in ("-h", "--help"):
        print(__doc__)
        return
    if args[0] == "help":
        print(HELP_EDITS)
        return
    fn = CMDS.get(args[0])
    if not fn:
        print(f"未知指令 {args[0]}\n{__doc__}")
        sys.exit(2)
    fn(args[1:])


if __name__ == "__main__":
    main()
